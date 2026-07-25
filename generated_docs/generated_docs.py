from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_document(title, sections, contents, filename):

    document = Document()

    # Page margins
    section = document.sections[0]
    section.top_margin = Pt(40)
    section.bottom_margin = Pt(40)
    section.left_margin = Pt(50)
    section.right_margin = Pt(40)

    # Title
    heading = document.add_heading()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(22)

    document.add_paragraph()

    # Section headings and content
    for section_name, content in zip(sections, contents):

        heading = document.add_heading(level=1)

        run = heading.add_run(section_name)
        run.bold = True
        run.font.size = Pt(16)

        para = document.add_paragraph(content)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        para.style.font.size = Pt(12)

        document.add_paragraph()

    document.save(filename)